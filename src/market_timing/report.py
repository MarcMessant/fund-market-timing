"""Export result tables: one CSV per table and one Excel workbook (one sheet per table).

Only estimates and aggregated statistics are written, never the underlying price series.
"""
from __future__ import annotations

import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = {
    "descriptive": "Descriptive statistics, weekly, CAPM",
    "treynor_mazuy": "Treynor-Mazuy (1966), weekly Wed-Wed, Newey-West",
    "treynor_mazuy_asym": "Treynor-Mazuy with separate up/down convexity",
    "henriksson_merton": "Henriksson-Merton (1981), weekly Wed-Wed, Newey-West",
    "april_2025": "April 2025 tariff crash, daily, implied betas",
    "anchor_day": "Robustness: week-end day (Mon to Fri)",
    "influence": "Robustness: influence of the week to 9 Apr 2025 (DFBETAS)",
    "dimson": "Robustness: Dimson (1979) lagged betas, weekly and daily",
    "monthly": "Robustness: monthly frequency (24 months, anchored on the 22nd)",
    "fama": "Fama (1972) decomposition, weekly, target beta = 1",
    "fama_monthly": "Fama (1972) decomposition, monthly, target beta = 1",
    "fama_pre_window_beta": "Fama with target beta = pre-window beta (discarded variant)",
    "bridge_hm": "Fama to timing bridge, Henriksson-Merton (exact OLS identity)",
    "bridge_tm": "Fama to timing bridge, Treynor-Mazuy (exact OLS identity)",
}

HEAD = PatternFill("solid", fgColor="1F2A44")


def export(res: dict, out: str) -> None:
    tables = os.path.join(out, "tables")
    os.makedirs(tables, exist_ok=True)
    for k in SHEETS:
        res[k].to_csv(os.path.join(tables, f"{k}.csv"), float_format="%.6f")

    path = os.path.join(out, "market_timing_results.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for k, title in SHEETS.items():
            df = res[k]
            df.to_excel(xw, sheet_name=k[:31], startrow=2)
            ws = xw.sheets[k[:31]]
            ws["A1"] = title
            ws["A1"].font = Font(bold=True, size=13)
            for c in ws[3]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEAD
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            for row in ws.iter_rows(min_row=4):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0000"
            ws.column_dimensions["A"].width = 22
            for i in range(2, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = 14
            ws.freeze_panes = "B4"
